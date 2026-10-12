from __future__ import annotations

from collections.abc import Mapping, Sequence
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal
from io import BytesIO
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from purchase_price.services.price_metrics import PriceResult

# Runtime marker: the export can carry the price metric contract (docs/metrics/price.md §6 차이 6).
MARKET_SURVEY_PRICE_CONTRACT_V1 = True
MARKET_SURVEY_EXPORT_SCHEMA = "workspace-v3-market-survey-xlsx-v2"
PRICE_BASIS_COLUMN = "가격 계산 포함"
DEFAULT_CAUTION = (
    "본 자료는 구매검토 보조자료입니다. 가격·단위·VAT·설치·운송·옵션·보증 등 "
    "조건 동일성을 원문에서 확인한 뒤 사용하세요. 시스템이 구매 여부나 업체를 결정하지 않습니다."
)


def _excel_value(value: Any) -> Any:
    if value is None:
        return "미확인"
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def _won(value: Decimal) -> str:
    return f"{value.quantize(Decimal(1), rounding=ROUND_HALF_UP):,}원"


def _signed_won(value: Decimal) -> str:
    sign = "+" if value > 0 else "-" if value < 0 else ""
    return f"{sign}{_won(abs(value))}"


def _price_summary_rows(result: PriceResult) -> list[tuple[str, str]]:
    summary = result.summary
    priced = f"{summary.priced_count}건"
    if summary.excluded_count > 0:
        units = ", ".join(summary.other_units) if summary.other_units else "없음"
        priced += (
            f" (제외 {summary.excluded_count}건: 입력 오류 의심 {summary.entry_error_count}"
            f" · 다른 단위 {summary.other_unit_count}({units}))"
        )
    if summary.priced_count == 0 or summary.median is None:
        median = "계산 안 함"
    else:
        median = _won(summary.median)
    if summary.low is not None and summary.high is not None:
        spread = f"{_won(summary.low)} ~ {_won(summary.high)}"
    else:
        spread = "계산 안 함"
    comparison = result.comparison
    rows = [
        ("거래 기간", result.period_label),
        ("기준 단위", summary.main_unit or "확인 안 됨"),
        ("확인 거래", f"{summary.confirmed_count}건"),
        ("가격 계산 거래", priced),
        ("가운데 값", median),
        ("최저~최고", spread),
    ]
    if summary.total_quantity is not None:
        rows.append(("총 수량", f"{summary.total_quantity.normalize():f}"))
    rows.extend(
        [
            ("구매 기관 수", f"{summary.institution_count}곳"),
            ("납품업체 수", f"{summary.supplier_count}곳"),
            ("최근 거래일", summary.latest_date or "미확인"),
            ("견적 비교 결과", comparison.label),
        ]
    )
    if comparison.difference_amount is not None and comparison.difference_percent is not None:
        percent = comparison.difference_percent
        percent_sign = "+" if percent > 0 else ""
        rows.append(
            (
                "견적 차이",
                f"{_signed_won(comparison.difference_amount)} ({percent_sign}{percent:.1f}%)",
            )
        )
    evidence = result.evidence
    evidence_text = evidence.label
    if evidence.reasons:
        evidence_text += " · " + ", ".join(evidence.reasons)
    rows.append(("근거 수준", evidence_text))
    as_of = result.data_as_of or "미확인"
    if result.data_as_of_stale:
        as_of += " (3일 넘게 지남)"
    rows.append(("자료 기준일", as_of))
    return rows


def _ordered_columns(rows: Sequence[Mapping[str, Any]]) -> list[str]:
    columns: list[str] = []
    seen: set[str] = set()
    for row in rows:
        for key in row:
            key_text = str(key)
            if key_text not in seen:
                seen.add(key_text)
                columns.append(key_text)
    return columns


def _append_table_sheet(
    workbook: Workbook,
    *,
    title: str,
    rows: Sequence[Mapping[str, Any]],
    empty_message: str,
) -> None:
    sheet = workbook.create_sheet(title=title)
    if not rows:
        sheet["A1"] = empty_message
        sheet["A1"].alignment = Alignment(wrap_text=True, vertical="top")
        sheet.column_dimensions["A"].width = 80
        return

    columns = _ordered_columns(rows)
    for column_index, name in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=column_index, value=name)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    for row_index, row in enumerate(rows, start=2):
        for column_index, name in enumerate(columns, start=1):
            sheet.cell(
                row=row_index,
                column=column_index,
                value=_excel_value(row.get(name)),
            ).alignment = Alignment(wrap_text=True, vertical="top")

    for column_index, name in enumerate(columns, start=1):
        longest = max(
            [len(str(name))]
            + [len(str(_excel_value(row.get(name)))) for row in rows]
        )
        sheet.column_dimensions[get_column_letter(column_index)].width = min(
            max(longest + 2, 12), 48
        )
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def build_market_survey_workbook(
    *,
    search_identity: Mapping[str, Any],
    quote_context: Mapping[str, Any] | None = None,
    direct_rows: Sequence[Mapping[str, Any]] = (),
    research_rows: Sequence[Mapping[str, Any]] = (),
    supplier_rows: Sequence[Mapping[str, Any]] = (),
    identity_rows: Sequence[Mapping[str, Any]] = (),
    data_as_of: str | None = None,
    caution: str = DEFAULT_CAUTION,
    generated_at: datetime | None = None,
    price_result: PriceResult | None = None,
    direct_row_basis: Sequence[str] | None = None,
) -> bytes:
    """Build the Workspace V3 market-survey Excel attachment.

    data_as_of is deliberately caller-supplied. When the connected source does not expose
    a trustworthy as-of timestamp, callers should pass None rather than infer one from
    transaction dates or export time.

    price_result, when given, adds the contract's price statistics to the summary sheet
    (never recomputed here). direct_row_basis is filled by the caller, one entry per
    direct_rows item ('포함' / '제외 (다른 단위)' / '제외 (입력 오류 의심)'); the extra
    A_B column appears only when price_result is given and the lengths match.
    """

    generated_at = generated_at or datetime.now(ZoneInfo("Asia/Seoul"))
    workbook = Workbook()
    summary = workbook.active
    summary.title = "시장조사 요약"

    summary_rows = [
        ("Export schema", MARKET_SURVEY_EXPORT_SCHEMA),
        ("보고서 생성일시", generated_at.isoformat()),
        ("data_as_of", data_as_of or "미확인"),
        ("주의", caution),
    ]
    summary_rows.extend(
        (str(key), _excel_value(value)) for key, value in search_identity.items()
    )
    if quote_context:
        summary_rows.extend(
            (f"견적.{key}", _excel_value(value))
            for key, value in quote_context.items()
        )

    if price_result is not None:
        summary_rows.extend(_price_summary_rows(price_result))

    if price_result is not None and direct_row_basis is not None and len(direct_row_basis) == len(direct_rows):
        direct_rows = [
            {**row, PRICE_BASIS_COLUMN: basis}
            for row, basis in zip(direct_rows, direct_row_basis)
        ]

    for row_index, (label, value) in enumerate(summary_rows, start=1):
        summary.cell(row=row_index, column=1, value=label).font = Font(bold=True)
        summary.cell(row=row_index, column=2, value=value).alignment = Alignment(
            wrap_text=True,
            vertical="top",
        )
    summary.column_dimensions["A"].width = 24
    summary.column_dimensions["B"].width = 92

    _append_table_sheet(
        workbook,
        title="A_B 직접근거",
        rows=direct_rows,
        empty_message="A/B 직접 비교 가능한 가격근거가 없습니다.",
    )
    _append_table_sheet(
        workbook,
        title="C_Research",
        rows=research_rows,
        empty_message="C/Research 참고근거가 없습니다.",
    )
    _append_table_sheet(
        workbook,
        title="조달업체",
        rows=supplier_rows,
        empty_message="A/B 직접근거 기준 조달 납품업체가 없습니다.",
    )
    _append_table_sheet(
        workbook,
        title="식약처 Identity",
        rows=identity_rows,
        empty_message="연결된 식약처 Identity 근거가 없습니다.",
    )

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
