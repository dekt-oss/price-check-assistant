"""병원 경영·동향 리포트 (Phase 4 자동 리포트).

One report = one hospital, one fiscal year, one peer group. Everything in it comes from the same
computed objects the Benchmark screen shows (``hospital_benchmark`` / ``hospital_metrics``), plus
optionally the News Radar titles and links of the last days (display-only, never AI input) and an
AI paragraph produced by ``hospital_ai_explanation`` from computed numbers only.

Two renderings: Markdown (for the weekly job summary, R2 and messenger) and an Excel workbook
(for the download button and for staff who keep working in Excel).
"""

from __future__ import annotations

import io
from collections.abc import Iterable, Sequence
from dataclasses import dataclass, field
from datetime import datetime, timedelta
from decimal import Decimal

from purchase_price.services import hospital_benchmark as benchmark
from purchase_price.services import hospital_master as master_service
from purchase_price.services import hospital_metrics as metrics
from purchase_price.services import news_radar as radar

HEADLINE_METRICS: tuple[str, ...] = (
    "revenue_growth",
    "medical_margin",
    "net_margin",
    "labor_ratio",
    "material_ratio",
    "admin_ratio",
    "revenue_per_bed",
)
TREND_METRICS: tuple[tuple[str, str, Decimal], ...] = (
    ("medical_revenue", "의료수익 (억원)", Decimal(100_000_000)),
    ("labor_ratio", "인건비율 (%)", Decimal(1)),
    ("material_ratio", "재료비율 (%)", Decimal(1)),
)
TREND_YEARS = 5
SOURCE_NAME = "의료기관 회계정보 공시(한국보건산업진흥원, haspa.khidi.or.kr)"
NEWS_GROUPS: tuple[str, ...] = ("our_hospital", "peer_hospitals")


def quality_notes(quality: benchmark.QualityReport) -> list[str]:
    """Plain sentences for the data-quality box, the AI input and the report."""

    notes: list[str] = []
    if quality.period_mismatch:
        notes.append(
            "회계기간이 기준 병원과 다른 병원이 있어 같은 연도라도 단순 비교에 주의해야 합니다: "
            + "; ".join(quality.period_mismatch)
        )
    if quality.no_data:
        notes.append("이 연도 회계자료가 없는 병원: " + ", ".join(quality.no_data))
    for name, accounts in quality.missing_accounts.items():
        notes.append(f"{name}은(는) 공시에 없는 항목이 있습니다: " + ", ".join(accounts))
    if quality.negative_equity:
        notes.append(
            "자본총계가 0 이하라 부채비율을 계산하지 않은 병원: " + ", ".join(quality.negative_equity)
        )
    if any("자료 없음" in text for text in quality.bed_counts.values()):
        notes.append("병상수 자료가 없는 병원은 병상당 지표와 유사 규모 비교를 할 수 없습니다.")
    if any("회계공시 일반현황" in text for text in quality.bed_counts.values()):
        notes.append("병상수는 회계공시 일반현황(그해 말 심평원 자료)의 값입니다. 심평원 병원정보 연계 후 확정값으로 바뀝니다.")
    if quality.unverified_types:
        notes.append("종별은 심평원 연계 전 1차 입력입니다: " + ", ".join(quality.unverified_types))
    return notes


@dataclass(frozen=True)
class NewsLine:
    published: str
    keywords: str
    title: str
    source: str
    url: str


@dataclass
class BenchmarkReport:
    target: master_service.Hospital
    fiscal_year: int
    peer_label: str
    peers: tuple[master_service.Hospital, ...]
    rows: list[metrics.ComparisonRow]
    findings: list[metrics.Finding]
    quality: list[str]
    period: tuple[str, str] | None
    full_table: list[dict[str, str]]
    trends: dict[str, dict[str, dict[int, Decimal | None]]]
    fetched: list[str]
    generated_at: datetime
    explanation: str | None = None
    explanation_note: str = ""
    news: list[NewsLine] = field(default_factory=list)


def build_report(
    data: benchmark.BenchmarkData,
    target: master_service.Hospital,
    peer_kind: str,
    fiscal_year: int,
    *,
    custom_ids: Sequence[str] = (),
    bed_range: tuple[int, int] | None = None,
    now: datetime | None = None,
) -> BenchmarkReport:
    peers = tuple(data.master.peer_group(target, peer_kind, custom_ids=custom_ids, bed_range=bed_range))
    rows = benchmark.compare(data, target, peers, fiscal_year, HEADLINE_METRICS)
    findings = (
        metrics.describe_findings(
            rows, benchmark.gap_history(data, target, peers, fiscal_year, HEADLINE_METRICS)
        )
        if peers
        else []
    )
    quality = benchmark.quality_report(data, target, peers, fiscal_year)
    hospitals = [target, *peers]
    full_table: list[dict[str, str]] = []
    for key, (label, unit, _) in metrics.METRIC_SPECS.items():
        entry = {"지표": label}
        for hospital in hospitals:
            values = benchmark.metrics_for(data, hospital, fiscal_year)
            entry[hospital.short_name] = metrics.format_metric(values.get(key) if values else None, unit)
        full_table.append(entry)
    years = list(range(fiscal_year - TREND_YEARS + 1, fiscal_year + 1))
    trends = {
        title: benchmark.trend_table(data, hospitals, key, years) for key, title, _ in TREND_METRICS
    }
    peer_label = master_service.PEER_GROUP_LABELS.get(peer_kind, peer_kind)
    if peer_kind == master_service.PEER_SIMILAR_SIZE and bed_range is not None:
        peer_label += f" {bed_range[0]}~{bed_range[1]}병상"
    return BenchmarkReport(
        target=target,
        fiscal_year=fiscal_year,
        peer_label=peer_label,
        peers=peers,
        rows=rows,
        findings=findings,
        quality=quality_notes(quality),
        period=benchmark.target_period(data, target.hospital_id, fiscal_year),
        full_table=full_table,
        trends=trends,
        fetched=sorted(set(quality.fetched.values())),
        generated_at=now or datetime.now().astimezone(radar.SEOUL),
    )


def news_lines(
    entries: Iterable[radar.NewsEntry],
    keyword_texts: Iterable[str],
    *,
    now: datetime,
    days: int = 7,
    limit: int = 20,
) -> list[NewsLine]:
    """Titles and links only (NAVER display terms), newest first, for the chosen keywords."""

    wanted = set(keyword_texts)
    since = now - timedelta(days=days)
    picked: list[NewsLine] = []
    for entry in radar.sorted_entries(entries):
        when = entry.published_at or entry.detected_at
        if when.tzinfo is None:
            when = when.astimezone()
        if when < since or not (set(entry.keywords) & wanted):
            continue
        picked.append(
            NewsLine(
                published=radar.seoul_time_text(entry.published_at or entry.detected_at),
                keywords=" · ".join(k for k in entry.keywords if k in wanted),
                title=entry.title,
                source=entry.source_domain or "출처 확인 안 됨",
                url=entry.url,
            )
        )
        if len(picked) >= limit:
            break
    return picked


def keywords_in_groups(group_keys: Iterable[str] = NEWS_GROUPS) -> list[str]:
    wanted = set(group_keys)
    return [kw.text for group in radar.load_keyword_groups() if group.key in wanted for kw in group.keywords]


def _trend_cell(value: Decimal | None, scale: Decimal) -> str:
    if value is None:
        return "-"
    return f"{value / scale:,.1f}"


def report_markdown(report: BenchmarkReport) -> str:
    t = report.target
    lines = [
        f"# {t.short_name} 경영 비교 리포트 ({report.fiscal_year} 회계연도)",
        "",
        f"- 만든 시각: {report.generated_at.astimezone(radar.SEOUL):%Y-%m-%d %H:%M} (한국 시간)",
        f"- 비교 방식: {report.peer_label} — " + (", ".join(p.short_name for p in report.peers) or "비교 병원 없음"),
        "- 회계기간: " + (f"{report.period[0]} ~ {report.period[1]}" if report.period else "확인 안 됨"),
        f"- 출처: {SOURCE_NAME}. 가져온 날: " + (", ".join(report.fetched) or "-"),
        "",
        "## 주요 경영지표",
        "",
        f"| 지표 | {t.short_name} | 비교군 평균 | 비교 병원 수 | 위치 |",
        "| --- | ---: | ---: | ---: | --- |",
    ]
    for row in report.rows:
        lines.append(
            f"| {row.label} | {metrics.format_metric(row.value, row.unit)} | "
            f"{metrics.format_metric(row.peer_average, row.unit)} | {row.peer_count}곳 | {row.position_text} |"
        )
    lines += ["", "## 해석", ""]
    if report.explanation:
        lines += [report.explanation, "", f"_{report.explanation_note}_", ""]
    if report.findings:
        lines += [f"- {f.text}" for f in report.findings]
    else:
        lines.append("- 비교군 평균과 눈에 띄게 다른 지표가 없거나, 비교할 자료가 부족합니다.")
    lines += ["", "숫자와 계산 문장은 모두 프로그램이 공시 원자료로 계산했습니다.", ""]
    lines += ["## 최근 5년 추이", ""]
    for key, title, scale in TREND_METRICS:
        table = report.trends.get(title, {})
        years = sorted({y for values in table.values() for y in values})
        lines += [f"**{title}**", "", "| 병원 | " + " | ".join(str(y) for y in years) + " |"]
        lines.append("| --- | " + " | ".join("---:" for _ in years) + " |")
        for name, values in table.items():
            lines.append(f"| {name} | " + " | ".join(_trend_cell(values.get(y), scale) for y in years) + " |")
        lines.append("")
    lines += ["## 자료 상태", ""]
    lines += [f"- {note}" for note in report.quality] or ["- 확인된 문제 없음"]
    if report.news:
        lines += ["", "## 최근 7일 관련 기사 (제목·링크만)", ""]
        for item in report.news:
            title = item.title.replace("[", "(").replace("]", ")")
            lines.append(f"- {item.published} · {item.keywords} · [{title}]({item.url}) · {item.source}")
    return "\n".join(lines).rstrip() + "\n"


def report_workbook(report: BenchmarkReport) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="E8EEF7")

    def write_table(ws, start_row: int, headers: Sequence[str], rows: Iterable[Sequence[object]]) -> int:
        for col, value in enumerate(headers, start=1):
            cell = ws.cell(row=start_row, column=col, value=value)
            cell.font = header_font
            cell.fill = header_fill
        r = start_row
        for r, values in enumerate(rows, start=start_row + 1):
            for col, value in enumerate(values, start=1):
                ws.cell(row=r, column=col, value=value)
        return r

    def autosize(ws, widths: dict[int, int] | None = None) -> None:
        for col in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(col)].width = (widths or {}).get(col, 16)

    t = report.target
    ws = wb.active
    ws.title = "요약"
    ws["A1"] = f"{t.short_name} 경영 비교 리포트 ({report.fiscal_year} 회계연도)"
    ws["A1"].font = Font(bold=True, size=14)
    meta = [
        ("만든 시각", f"{report.generated_at.astimezone(radar.SEOUL):%Y-%m-%d %H:%M} (한국 시간)"),
        ("비교 방식", f"{report.peer_label}: " + (", ".join(p.short_name for p in report.peers) or "없음")),
        ("회계기간", f"{report.period[0]} ~ {report.period[1]}" if report.period else "확인 안 됨"),
        ("출처", f"{SOURCE_NAME} / 가져온 날 {', '.join(report.fetched) or '-'}"),
    ]
    for i, (k, v) in enumerate(meta, start=2):
        ws.cell(row=i, column=1, value=k).font = header_font
        ws.cell(row=i, column=2, value=v)
    end = write_table(
        ws,
        7,
        ["지표", t.short_name, "비교군 평균", "비교 병원 수", "위치"],
        (
            [
                row.label,
                metrics.format_metric(row.value, row.unit),
                metrics.format_metric(row.peer_average, row.unit),
                f"{row.peer_count}곳",
                row.position_text,
            ]
            for row in report.rows
        ),
    )
    r = end + 2
    ws.cell(row=r, column=1, value="해석").font = header_font
    r += 1
    if report.explanation:
        cell = ws.cell(row=r, column=1, value=report.explanation)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
        ws.row_dimensions[r].height = 120
        r += 1
        ws.cell(row=r, column=1, value=report.explanation_note)
        r += 1
    for finding in report.findings or []:
        ws.cell(row=r, column=1, value=f"- {finding.text}")
        r += 1
    if not report.findings:
        ws.cell(row=r, column=1, value="- 비교군 평균과 눈에 띄게 다른 지표가 없거나, 비교할 자료가 부족합니다.")
        r += 1
    ws.cell(row=r + 1, column=1, value="숫자와 계산 문장은 모두 프로그램이 공시 원자료로 계산했습니다.")
    autosize(ws, {1: 26, 2: 40})

    ws2 = wb.create_sheet("병원별 전체 지표")
    headers = list(report.full_table[0].keys()) if report.full_table else ["지표"]
    write_table(ws2, 1, headers, ([row.get(h, "") for h in headers] for row in report.full_table))
    autosize(ws2, {1: 26})

    ws3 = wb.create_sheet("5년 추이")
    r = 1
    for key, title, scale in TREND_METRICS:
        table = report.trends.get(title, {})
        years = sorted({y for values in table.values() for y in values})
        ws3.cell(row=r, column=1, value=title).font = Font(bold=True, size=12)
        r = write_table(
            ws3,
            r + 1,
            ["병원", *[str(y) for y in years]],
            (
                [name, *[(float(v / scale) if (v := values.get(y)) is not None else None) for y in years]]
                for name, values in table.items()
            ),
        ) + 2
    autosize(ws3, {1: 22})

    ws4 = wb.create_sheet("자료 상태")
    ws4["A1"] = "자료 상태"
    ws4["A1"].font = Font(bold=True, size=12)
    for i, note in enumerate(report.quality or ["확인된 문제 없음"], start=2):
        ws4.cell(row=i, column=1, value=note)
    autosize(ws4, {1: 120})

    if report.news:
        ws5 = wb.create_sheet("관련 기사")
        write_table(
            ws5,
            1,
            ["게시 시각", "키워드", "제목", "출처", "링크"],
            ([n.published, n.keywords, n.title, n.source, n.url] for n in report.news),
        )
        autosize(ws5, {3: 70, 5: 60})

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def report_filename(report: BenchmarkReport, extension: str) -> str:
    return f"{report.target.short_name}_경영비교_{report.fiscal_year}_{report.peer_label}.{extension}".replace(" ", "")
