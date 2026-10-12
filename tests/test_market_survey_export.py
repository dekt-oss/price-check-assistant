from __future__ import annotations

import json
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

from purchase_price.services import price_metrics
from purchase_price.services.market_survey_export import (
    MARKET_SURVEY_EXPORT_SCHEMA,
    MARKET_SURVEY_PRICE_CONTRACT_V1,
    build_market_survey_workbook,
)

FIXTURES = Path(__file__).parent / "fixtures" / "price"
GENERATED_AT = datetime(2026, 10, 12, 9, 0, tzinfo=ZoneInfo("Asia/Seoul"))


def _sheet_rows(workbook, title: str) -> list[list[object]]:
    sheet = workbook[title]
    return [list(row) for row in sheet.iter_rows(values_only=True)]


def test_market_survey_export_preserves_direct_research_source_and_identity() -> None:
    payload = build_market_survey_workbook(
        search_identity={
            "검색어": "C101",
            "품목": "채혈기",
            "모델": "C101",
            "식약처 품목번호": "수신 22-2177호",
        },
        quote_context={
            "내 견적가": 500,
            "단위": "개",
            "VAT": "미확인",
            "설치·운송 등 조건": "미확인",
        },
        direct_rows=[
            {
                "가격": "450원",
                "단위": "개",
                "매칭등급": "A",
                "매칭근거": "model_exact_and_product_match",
                "판매처": "공급사A",
                "구매처": "기관A",
                "Source": "나라장터 납품요구",
                "원문근거키": "raw/example.json.gz",
            }
        ],
        research_rows=[
            {
                "가격": "470원",
                "자료성격": "검색 참고",
                "Source": "공개조달 Research",
                "원문": "https://example.invalid/source",
            }
        ],
        supplier_rows=[
            {
                "공급업체": "공급사A",
                "직접거래건수": 1,
                "Source": "나라장터 납품요구",
            }
        ],
        identity_rows=[
            {
                "유형": "신고",
                "식약처 품목번호": "수신 22-2177호",
                "모델": "C101",
                "품목 책임주체": "책임주체A",
                "UDI-DI": "UDI-1",
                "Source": "식약처 제품정보",
            }
        ],
        data_as_of=None,
        generated_at=datetime(2026, 9, 28, 15, 0, tzinfo=ZoneInfo("Asia/Seoul")),
    )

    workbook = load_workbook(BytesIO(payload), data_only=True)

    assert workbook.sheetnames == [
        "시장조사 요약",
        "A_B 직접근거",
        "C_Research",
        "조달업체",
        "식약처 Identity",
    ]
    summary = dict(_sheet_rows(workbook, "시장조사 요약"))
    assert summary["Export schema"] == MARKET_SURVEY_EXPORT_SCHEMA
    assert summary["data_as_of"] == "미확인"
    assert "구매검토 보조자료" in summary["주의"]
    assert summary["검색어"] == "C101"
    assert summary["견적.내 견적가"] == 500

    direct = _sheet_rows(workbook, "A_B 직접근거")
    direct_header = direct[0]
    direct_values = dict(zip(direct_header, direct[1], strict=True))
    assert direct_values["매칭등급"] == "A"
    assert direct_values["Source"] == "나라장터 납품요구"
    assert direct_values["원문근거키"] == "raw/example.json.gz"

    research = _sheet_rows(workbook, "C_Research")
    research_values = dict(zip(research[0], research[1], strict=True))
    assert research_values["자료성격"] == "검색 참고"
    assert research_values["원문"] == "https://example.invalid/source"

    identity = _sheet_rows(workbook, "식약처 Identity")
    identity_values = dict(zip(identity[0], identity[1], strict=True))
    assert identity_values["식약처 품목번호"] == "수신 22-2177호"
    assert identity_values["UDI-DI"] == "UDI-1"


def test_market_survey_export_marks_empty_sections_without_inventing_evidence() -> None:
    payload = build_market_survey_workbook(
        search_identity={"검색어": "UNKNOWN"},
        generated_at=datetime(2026, 9, 28, 15, 0, tzinfo=ZoneInfo("Asia/Seoul")),
    )
    workbook = load_workbook(BytesIO(payload), data_only=True)

    assert workbook["A_B 직접근거"]["A1"].value == "A/B 직접 비교 가능한 가격근거가 없습니다."
    assert workbook["C_Research"]["A1"].value == "C/Research 참고근거가 없습니다."
    assert workbook["조달업체"]["A1"].value == "A/B 직접근거 기준 조달 납품업체가 없습니다."
    assert workbook["식약처 Identity"]["A1"].value == "연결된 식약처 Identity 근거가 없습니다."


def _price_result(name: str, *, quote: str) -> price_metrics.PriceResult:
    data = json.loads((FIXTURES / f"{name}.json").read_text(encoding="utf-8"))
    trades = [price_metrics.PriceTrade.from_mapping(row) for row in data["trades"]]
    valid = [trade for trade in trades if trade.entry_error_reason is None]
    return price_metrics.build_price_result(
        trades,
        quote_unit_price=Decimal(quote),
        identity_confirmed=True,
        identity_exact=True,
        period_label="최근 3년",
        cutoff=date(2023, 10, 13),
        preferred_unit=price_metrics.main_unit(valid),
        data_as_of="2026-10-10",
        today=date(2026, 10, 12),
    )


def _summary(payload: bytes) -> dict[str, object]:
    return dict(_sheet_rows(load_workbook(BytesIO(payload), data_only=True), "시장조사 요약"))


def test_price_contract_marker_and_schema_version() -> None:
    assert MARKET_SURVEY_PRICE_CONTRACT_V1 is True
    assert MARKET_SURVEY_EXPORT_SCHEMA.endswith("-v2")


def test_market_survey_export_adds_price_contract_rows_for_mixed_units() -> None:
    result = _price_result("many_mixed_units", quote="6500000")
    payload = build_market_survey_workbook(
        search_identity={"검색어": "M40"},
        price_result=result,
        generated_at=GENERATED_AT,
    )
    summary = _summary(payload)

    assert summary["거래 기간"] == "최근 3년"
    assert summary["기준 단위"] == "대"
    assert summary["확인 거래"] == "18건"
    assert summary["가격 계산 거래"].startswith("16건")
    assert "제외 2건" in summary["가격 계산 거래"]
    assert "다른 단위 2(" in summary["가격 계산 거래"]
    assert summary["가운데 값"] == "4,400,000원"
    assert summary["최저~최고"] == "4,389,000원 ~ 6,585,500원"
    assert summary["총 수량"] == "16"
    assert summary["구매 기관 수"] == "14곳"
    assert summary["납품업체 수"] == "7곳"
    assert summary["최근 거래일"] == "2026-09-01"
    assert summary["견적 비교 결과"] == "확인 필요"
    assert summary["견적 차이"] == "+2,100,000원 (+47.7%)"
    assert summary["근거 수준"].startswith("보통")
    assert summary["자료 기준일"] == "2026-10-10"


def test_market_survey_export_marks_stale_data_as_of() -> None:
    result = _price_result("many_mixed_units", quote="6500000")
    stale = price_metrics.PriceResult(
        summary=result.summary,
        comparison=result.comparison,
        evidence=result.evidence,
        period_label=result.period_label,
        data_as_of="2026-10-01",
        data_as_of_stale=True,
        show_category_in_conclusion=result.show_category_in_conclusion,
    )
    summary = _summary(
        build_market_survey_workbook(search_identity={}, price_result=stale, generated_at=GENERATED_AT)
    )
    assert summary["자료 기준일"] == "2026-10-01 (3일 넘게 지남)"


def test_market_survey_export_says_not_calculated_when_no_priced_trades() -> None:
    result = _price_result("zero_trades", quote="6500000")
    assert result.summary.priced_count == 0
    summary = _summary(
        build_market_survey_workbook(
            search_identity={"검색어": "ZERO"}, price_result=result, generated_at=GENERATED_AT
        )
    )

    assert summary["가운데 값"] == "계산 안 함"
    assert summary["최저~최고"] == "계산 안 함"
    assert summary["기준 단위"] == "확인 안 됨"
    assert summary["가격 계산 거래"] == "0건"
    assert "견적 차이" not in summary


def test_market_survey_export_without_price_result_has_no_price_rows() -> None:
    payload = build_market_survey_workbook(
        search_identity={"검색어": "M40"}, generated_at=GENERATED_AT
    )
    summary = _summary(payload)
    assert "가운데 값" not in summary
    assert "근거 수준" not in summary
    assert list(summary)[:4] == ["Export schema", "보고서 생성일시", "data_as_of", "주의"]


def test_direct_rows_get_price_basis_column_only_when_caller_supplies_it() -> None:
    result = _price_result("many_mixed_units", quote="6500000")
    rows = [{"가격": "4,400,000원", "단위": "대"}, {"가격": "1,000원", "단위": "set"}]

    def header(**kwargs) -> list[object]:
        payload = build_market_survey_workbook(
            search_identity={}, direct_rows=rows, generated_at=GENERATED_AT, **kwargs
        )
        return _sheet_rows(load_workbook(BytesIO(payload)), "A_B 직접근거")[0]

    assert "가격 계산 포함" not in header(price_result=result)
    assert "가격 계산 포함" not in header(price_result=result, direct_row_basis=["포함"])
    assert "가격 계산 포함" not in header(direct_row_basis=["포함", "제외 (다른 단위)"])

    payload = build_market_survey_workbook(
        search_identity={},
        direct_rows=rows,
        price_result=result,
        direct_row_basis=["포함", "제외 (다른 단위)"],
        generated_at=GENERATED_AT,
    )
    direct = _sheet_rows(load_workbook(BytesIO(payload)), "A_B 직접근거")
    assert direct[0][-1] == "가격 계산 포함"
    assert [row[-1] for row in direct[1:]] == ["포함", "제외 (다른 단위)"]
