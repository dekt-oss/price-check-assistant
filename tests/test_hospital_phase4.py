"""Phase 4: AI explanation guard, weekly report, and the report workflow contract."""

from __future__ import annotations

import io
from datetime import datetime, timedelta, timezone
from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace

import pytest

from purchase_price.services import hospital_ai_explanation as ai
from purchase_price.services import hospital_benchmark as benchmark
from purchase_price.services import hospital_metrics as metrics
from purchase_price.services import hospital_report as report_service
from purchase_price.services import news_radar as radar

ROOT = Path(__file__).resolve().parents[1]
KST = timezone(timedelta(hours=9))


def _rows() -> list[metrics.ComparisonRow]:
    return [
        metrics.ComparisonRow("labor_ratio", "인건비율", "%", Decimal("50.4"), Decimal("47.6"), 4, "above"),
        metrics.ComparisonRow("material_ratio", "재료비율", "%", Decimal("35.9"), Decimal("36.4"), 4, "average"),
    ]


def _source() -> ai.ExplanationInput:
    rows = _rows()
    return ai.build_input(
        target_name="부산백병원",
        fiscal_year=2024,
        peer_label="지역 경쟁군",
        peer_names=["부산대병원", "동아대병원"],
        rows=rows,
        findings=metrics.describe_findings(rows),
        quality_notes=["회계기간이 다른 병원: 부산대병원: 2024-01-01 ~ 2024-12-31"],
    )


class _FakeMessages:
    def __init__(self, text: str, stop_reason: str = "end_turn") -> None:
        self.text = text
        self.stop_reason = stop_reason
        self.calls: list[dict] = []

    def create(self, **kwargs):
        self.calls.append(kwargs)
        return SimpleNamespace(
            stop_reason=self.stop_reason,
            model=kwargs["model"],
            content=[
                SimpleNamespace(type="thinking", thinking=""),
                SimpleNamespace(type="text", text=self.text),
            ],
            usage=SimpleNamespace(input_tokens=900, output_tokens=150),
        )


def _client(text: str, stop_reason: str = "end_turn"):
    messages = _FakeMessages(text, stop_reason)
    return SimpleNamespace(beta=SimpleNamespace(messages=messages)), messages


def test_input_contains_only_computed_display_values() -> None:
    payload = _source().to_json()
    assert "50.4%" in payload and "47.6%" in payload
    assert "인건비율이 비교군 평균보다 2.8%p 높습니다." in payload
    assert payload == _source().to_json()  # byte-stable for prompt caching


def test_explanation_request_uses_current_model_fallbacks_and_cached_system() -> None:
    client, messages = _client(
        "2024 회계연도 부산백병원의 인건비율은 50.4%로 비교군 평균 47.6%보다 2.8%p 높습니다. "
        "부산대병원은 회계기간이 달라 주의가 필요합니다."
    )
    result = ai.explain(_source(), api_key=None, client=client)
    assert "50.4%" in result.text
    call = messages.calls[0]
    assert call["model"] == "claude-opus-5-5"
    assert call["fallbacks"] == "default"
    assert call["betas"] == ["server-side-fallback-2026-07-01"]
    assert call["system"][0]["cache_control"] == {"type": "ephemeral"}
    assert "thinking" not in call  # Opus 5.5: thinking is always on; effort controls depth
    assert call["output_config"] == {"effort": "low"}


def test_answer_with_an_invented_number_is_rejected() -> None:
    client, _ = _client("인건비율이 평균보다 3.5%p 높고 인건비가 120억원 늘었습니다.")
    with pytest.raises(ai.ExplanationError, match="계산 결과에 없는 숫자"):
        ai.explain(_source(), api_key=None, client=client)


def test_refusal_and_empty_answers_fall_back() -> None:
    client, _ = _client("", stop_reason="refusal")
    with pytest.raises(ai.ExplanationError):
        ai.explain(_source(), api_key=None, client=client)
    client, _ = _client("   ")
    with pytest.raises(ai.ExplanationError):
        ai.explain(_source(), api_key=None, client=client)


def test_missing_key_never_calls_the_api(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    assert ai.resolve_api_key({}) is None
    assert ai.resolve_api_key({"ANTHROPIC_API_KEY": "  k  "}) == "k"
    with pytest.raises(ai.ExplanationError, match="연결 설정"):
        ai.explain(_source(), api_key=None)


def test_unsupported_numbers_accepts_formatted_input_values() -> None:
    source = _source()
    assert ai.unsupported_numbers("인건비율 50.4%, 평균 47.6%, 차이 2.8%p (2024)", source) == set()
    assert ai.unsupported_numbers("차이가 3년째 커졌습니다", source) == {"3"}


def test_weekly_report_from_committed_data_has_real_rows_and_no_invented_values() -> None:
    data = benchmark.load_benchmark_data()
    target = data.master.get("H-BUSAN-PAIK")
    assert target is not None
    report = report_service.build_report(data, target, "region", 2024, now=datetime(2026, 10, 12, 8, tzinfo=KST))
    text = report_service.report_markdown(report)
    assert "| 인건비율 | 50.4% | 47.6% | 4곳 | 높은 편 |" in text
    assert "| 병상당 의료수익 | 자료 없음 | 자료 없음 | 0곳 | 비교 불가 |" in text
    assert "2024-03-01 ~ 2025-02-28" in text
    assert any("회계기간" in note for note in report.quality)
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(report_service.report_workbook(report)))
    assert wb.sheetnames[:4] == ["요약", "병원별 전체 지표", "5년 추이", "자료 상태"]
    assert wb["요약"]["B8"].value == "-14.1%"


def test_news_lines_keep_only_recent_titles_for_chosen_keywords() -> None:
    now = datetime(2026, 10, 12, 8, tzinfo=KST)

    def entry(url: str, keyword: str, days_ago: int) -> radar.NewsEntry:
        when = now - timedelta(days=days_ago)
        return radar.NewsEntry(
            article_id=url, title=f"기사 {url}", url=url, naver_link="", source_domain="e.kr",
            published_at=when, keywords=(keyword,), detected_at=when,
        )

    lines = report_service.news_lines(
        [entry("a", "부산백병원", 1), entry("b", "의료 AI 병원", 1), entry("c", "부산백병원", 9)],
        ["부산백병원"],
        now=now,
    )
    assert [line.url for line in lines] == ["a"]
    assert lines[0].published == "10-11 08:00"


def test_weekly_report_workflow_contract() -> None:
    text = (ROOT / ".github" / "workflows" / "hospital-weekly-report.yml").read_text(encoding="utf-8")
    assert 'cron: "0 23 * * 0"' in text  # Monday 08:00 KST
    for secret in ("R2_ACCESS_KEY_ID", "R2_SECRET_ACCESS_KEY", "ANTHROPIC_API_KEY"):
        assert f"secrets.{secret}" in text
    assert "--with-ai --with-news --r2" in text
    assert "retention-days: 21" in text
    assert "NAVER_CLIENT" not in text  # the report reads stored titles; it never searches NAVER


def test_benchmark_page_keeps_plain_words_and_ai_fallback_text() -> None:
    from purchase_price.ui import result_summary as rs

    source = (ROOT / "pages" / "21_병원_경영_Benchmark.py").read_text(encoding="utf-8")
    assert "AI 설명 보기" in source and "엑셀 리포트 내려받기" in source
    assert "위 계산 문장으로 확인해 주세요" in source
    assert rs.has_banned_term(source) is None
